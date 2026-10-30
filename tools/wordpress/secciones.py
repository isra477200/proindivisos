"""Secciones nuevas del 30/10/2026: impuestos por comunidad, modelos descargables, calculadoras 2 y preguntas frecuentes.
Uso: python3 secciones.py fiscal|modelos|herramientas|faq [--solo slug [slug...]] [--sin-imagenes] [--sin-yoastqa]
Por cada spec: imagen MiniMax (manifest_<sección>.json en el scratchpad, gen.py + upload.py) → página (elementor-create-page si no
existe) → padre = índice de la sección (si la sección tiene índice) → destacada → composición con pagebuilder.run → metas Yoast →
publicada. El índice se construye primero para que las hijas tengan padre.
Extras: modelos → genera el .docx de cada modelo (python-docx), lo sube a medios y guarda la opción zp_modelo_{slug} {titulo, texto,
docx} que pinta [zp_modelo] (snippet 226); faq → guarda zp_faq_bloque_{n} = [[pregunta, respuesta], ...] que pinta [zp_faq_bloque]
(snippet 226, con un único FAQPage). Los campos internos de los specs (datos, fuentes_verificacion, dudas, modelo, preguntas) no se publican."""
import sys, os, json, subprocess, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import blogbuilder as b
from pagebuilder import run, meta
HERE = os.path.dirname(os.path.abspath(__file__))
IMGDIR = os.environ.get('ZP_IMGDIR', '/tmp/claude-0/-home-claude/465a46cd-02da-53eb-bb5e-39c7569fb424/scratchpad/img')
CFG = {
    'fiscal': {'hub': 'impuestos-del-proindiviso', 'files': ['fiscal_hub.json', 'fiscal_1.json', 'fiscal_2.json', 'fiscal_3.json'], 'pref': 'impuestos-proindiviso-'},
    'modelos': {'hub': 'modelos-documentos-proindiviso', 'files': ['modelos.json'], 'pref': 'imagen-'},
    'herramientas': {'hub': None, 'files': ['herramientas.json'], 'pref': 'imagen-'},
    'faq': {'hub': None, 'files': ['faq.json'], 'pref': 'imagen-'},
}
env = {}
for l in open(os.path.join(HERE, 'wp.env')):
    if '=' in l: k, v = l.strip().split('=', 1); env[k] = v.strip()

def option(name, value):
    out = subprocess.run(['curl', '-sS', '--max-time', '60', '-u', env['P1'], '-X', 'POST', '-H', 'Content-Type: application/json',
                          env['WP'] + '/wp-json/zp/v1/option/' + name, '-d', json.dumps({'value': value}, ensure_ascii=False)], capture_output=True, text=True).stdout
    ok = out.strip().startswith('{') or out.strip().startswith('[')
    print('  opción', name, 'OK' if ok else 'FAIL ' + out[:160]); return ok

def docx_modelo(s):
    """Genera el .docx del modelo (Manrope no viaja en Word: Calibri), lo sube a medios y devuelve su URL."""
    from docx import Document
    from docx.shared import Pt
    m = s['modelo']; d = Document()
    st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
    d.add_heading(m['titulo'], level=1)
    for p in re.split(r'\n\s*\n', m['texto'].strip()):
        d.add_paragraph(p.strip())
    if m.get('instrucciones'):
        d.add_heading('Cómo rellenarlo', level=2)
        for x in m['instrucciones']: d.add_paragraph(x, style='List Bullet')
    d.add_paragraph('Fuente: zonaproindiviso.com — modelo orientativo; revísalo con un abogado si hay importes altos o conflicto.')
    fn = s['slug'] + '.docx'; path = os.path.join(IMGDIR, 'out', fn); d.save(path)
    r = b.api('GET', f"/wp/v2/media?search={s['slug']}&per_page=10&_fields=id,source_url,mime_type")
    for x in (r if isinstance(r, list) else []):
        if x.get('source_url', '').endswith('/' + fn):
            subprocess.run(['curl', '-sS', '-u', env['P1'], '-X', 'DELETE', env['WP'] + f"/wp-json/wp/v2/media/{x['id']}?force=true", '-o', '/dev/null'])
    out = subprocess.run(['curl', '-sS', '--max-time', '120', '-u', env['P1'], '-X', 'POST', env['WP'] + '/wp-json/wp/v2/media',
                          '-H', f'Content-Disposition: attachment; filename="{fn}"',
                          '-H', 'Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document', '--data-binary', '@' + path], capture_output=True, text=True).stdout
    try:
        dd = json.loads(out); b.api('POST', f"/wp/v2/media/{dd['id']}", {'title': m['titulo'], 'alt_text': m['titulo']}); return dd['source_url']
    except Exception:
        print('  DOCX UPLOAD FAIL', out[:200]); return ''

def main():
    args = sys.argv[1:]; sec = args[0]; c = CFG[sec]
    solo = args[args.index('--solo') + 1:] if '--solo' in args else []
    solo = [x for x in solo if not x.startswith('--')]
    S = []
    for f in c['files']:
        p = os.path.join(HERE, '..', 'blog', f)
        if os.path.exists(p): S += json.load(open(p))
        else: print('falta', f)
    if c['hub']: S.sort(key=lambda s: 0 if s['slug'] == c['hub'] else 1)
    if solo: S = [s for s in S if s['slug'] in solo or (c['hub'] and s['slug'] == c['hub'] and c['hub'] in solo)]
    ids = b.load_ids(); IMG_F = os.path.join(HERE, 'images.json')
    MF = f'manifest_{sec}.json'; MAP_F = os.path.join(IMGDIR, MF.replace('manifest', 'media_map'))
    if '--sin-imagenes' not in args:
        man = json.load(open(os.path.join(IMGDIR, MF))) if os.path.exists(os.path.join(IMGDIR, MF)) else []
        have = {m['key'] for m in man}
        for s in S:
            if s['slug'] in have: continue
            man.append({'key': s['slug'], 'kind': 'page', 'post_id': 0, 'filename': (c['pref'] + s['slug'])[:70] + '.webp', 'title': s['h1'][:120], 'alt': s['imagen_alt'], 'prompt': s['imagen_prompt']})
        json.dump(man, open(os.path.join(IMGDIR, MF), 'w'), ensure_ascii=False, indent=1)
        keys = [s['slug'] for s in S]
        r = subprocess.run(['python3', 'gen.py'] + keys, cwd=IMGDIR, env={**os.environ, 'MANIFEST': MF}, capture_output=True, text=True); print(r.stdout[-400:])
        r = subprocess.run(['python3', 'upload.py'] + [k for k in keys if k not in json.load(open(MAP_F))] if os.path.exists(MAP_F) else ['python3', 'upload.py'] + keys,
                           cwd=IMGDIR, env={**os.environ, 'MANIFEST': MF}, capture_output=True, text=True); print(r.stdout[-400:])
    mm = json.load(open(MAP_F)) if os.path.exists(MAP_F) else {}
    hub_id = None; hechos = []
    for s in S:
        pid = b.ensure(s, ids)
        if not pid: print('SIN ID', s['slug']); continue
        es_hub = c['hub'] and s['slug'] == c['hub']
        if es_hub: hub_id = pid
        elif c['hub']:
            if hub_id is None: hub_id = ids.get(c['hub']) or b.ensure({'slug': c['hub'], 'kind': 'page', 'title_corto': c['hub']}, ids)
            b.api('POST', f'/wp/v2/pages/{pid}', {'parent': hub_id})
        mid = mm.get(s['slug'], {}).get('media_id')
        if mid:
            IMG = json.load(open(IMG_F)); IMG[str(pid)] = mid; json.dump(IMG, open(IMG_F, 'w'), indent=0)
            b.api('POST', f'/wp/v2/pages/{pid}', {'featured_media': mid})
        else: print('  sin imagen', s['slug'])
        if sec == 'modelos' and s.get('modelo'):
            url = docx_modelo(s)
            option('zp_modelo_' + s['slug'], {'titulo': s['modelo']['titulo'], 'texto': s['modelo']['texto'], 'docx': url})
        if sec == 'faq':
            n = 0
            for x in s['sections']:
                m = re.search(r'zp_faq_bloque n="(\d+)"', x.get('shortcode', ''))
                if m: option('zp_faq_bloque_' + m.group(1), [list(q) for q in x.get('preguntas', [])]); n = max(n, int(m.group(1)))
            for k in range(n + 1, 13): option('zp_faq_bloque_' + str(k), [])
        sp = b.normalize({k: v for k, v in s.items() if k not in ('datos', 'fuentes_verificacion', 'dudas', 'modelo')})
        sp['post_id'] = pid; sp['dated'] = sec in ('fiscal', 'herramientas')
        if mid: sp['image_id'] = mid
        ok = run(sp)
        if ok:
            meta(pid, title=s['title_corto'], yt=s['yoast_title'].strip(), yd=s['meta_desc'], kw=s['focus_kw'], status='publish',
                 parent=(None if es_hub or not c['hub'] else hub_id), ptype='pages', slug=s['slug'])
            hechos.append(s)
        print(pid, s['slug'], 'OK' if ok else 'FAIL', flush=True)
    if hechos and '--sin-yoastqa' not in args:
        urls = []
        for s in hechos:
            r = b.api('GET', f"/wp/v2/pages?slug={s['slug']}&_fields=link")
            if isinstance(r, list) and r: urls.append(r[0]['link'])
        r = subprocess.run(['python3', 'yoastqa.py', '--urls'] + urls + ['--save'], cwd=HERE, capture_output=True, text=True); print(r.stdout[-2500:])

if __name__ == '__main__':
    main()
